#!/usr/bin/env python3
"""S02 - Structural exploration of every scientific file group.

Reads (without pasting wholesale into any report):
  * the 5 normalized complete dynamic spectra
  * the 12 drift-population distributions in `drift_population_series.tsv`
  * the 6 directional count series in `orientation_count_series.tsv`
  * all 14 image sheets of the source workbook, to determine how each maps to
    the normalized views and whether any image sheet is unrepresented.
"""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
OUT = ROOT / "output" / "results"

np.set_printoptions(suppress=True, precision=6)


def load_workbook_grid(sheet: str) -> np.ndarray:
    """Return the 425x425 numeric image block of a workbook image sheet."""
    import openpyxl

    wb = openpyxl.load_workbook(
        DATA / "source_data" / "all_source_data_deidentified.xlsx",
        read_only=True,
        data_only=True,
    )
    ws = wb[sheet]
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=6, max_col=427, values_only=True)):
        vals = np.array(
            [np.nan if v is None else float(v) for v in list(row)[2:]], dtype=float
        )
        rows.append(vals)
        if len(rows) >= 425:
            break
    wb.close()
    return np.vstack(rows)


def main() -> None:
    report = {}

    # ---------- 1. normalized dynamic spectra ----------
    meta = pd.read_csv(DATA / "normalized_views" / "dynamic_spectrum_event_metadata.tsv", sep="\t")
    ev = {}
    for _, m in meta.iterrows():
        eid = m["event_id"]
        df = pd.read_csv(DATA / "normalized_views" / f"{eid}_dynamic_spectrum.tsv", sep="\t")
        times = np.array([float(c.replace("time_s_", "")) for c in df.columns[1:]])
        freq = df["frequency_MHz"].to_numpy(float)
        img = df.iloc[:, 1:].to_numpy(float)
        ev[eid] = dict(
            freq=freq, times=times, img=img,
            shape=list(img.shape),
            tmin=float(m["time_min_s"]), tmax=float(m["time_max_s"]),
            fmin=float(m["frequency_min_MHz"]), fmax=float(m["frequency_max_MHz"]),
            reported=float(m["reported_event_drift_rate_MHz_s"]),
        )
        print(
            f"{eid}: grid={img.shape} dt={np.diff(times).mean():.6f}s "
            f"df={np.diff(freq).mean():.6f}MHz  reported_drift={m['reported_event_drift_rate_MHz_s']:.4f} "
            f"img[min={img.min():.4f} max={img.max():.4f} med={np.median(img):.4f}]"
        )
    report["normalized_events"] = {
        k: {kk: vv for kk, vv in v.items() if kk not in ("freq", "times", "img")}
        for k, v in ev.items()
    }
    report["normalized_events"]["_time_span_check"] = {
        k: {
            "times_span_s": float(v["times"].max() - v["times"].min()),
            "freq_span_MHz": float(v["freq"].max() - v["freq"].min()),
            "n_t": int(v["times"].size),
            "n_f": int(v["freq"].size),
        }
        for k, v in ev.items()
    }

    # ---------- 2. drift population series ----------
    pop = pd.read_csv(DATA / "normalized_views" / "drift_population_series.tsv", sep="\t")
    print("\n=== drift_population_series ===")
    print(pop.groupby(["population_series_id", "anonymous_source_group", "polarization_channel", "drift_regime"]).agg(
        n_bins=("distribution_bin_center", "size"),
        bin_min=("distribution_bin_center", "min"),
        bin_max=("distribution_bin_center", "max"),
        events=("event_count", "sum"),
    ).to_string())
    report["population_series_table"] = (
        pop.groupby(["population_series_id", "anonymous_source_group", "polarization_channel", "drift_regime"])
        .agg(n_bins=("distribution_bin_center", "size"),
             bin_min=("distribution_bin_center", "min"),
             bin_max=("distribution_bin_center", "max"),
             events=("event_count", "sum"))
        .reset_index()
        .to_dict("records")
    )

    # ---------- 3. orientation count series ----------
    ori = pd.read_csv(DATA / "normalized_views" / "orientation_count_series.tsv", sep="\t")
    print("\n=== orientation_count_series ===")
    print(ori.groupby("orientation_series_id").agg(
        n=("relative_coordinate", "size"),
        cmin=("relative_coordinate", "min"),
        cmax=("relative_coordinate", "max"),
        events=("event_count", "sum"),
        zero_frac=("event_count", lambda s: float((s == 0).mean())),
    ).to_string())
    report["orientation_table"] = (
        ori.groupby("orientation_series_id")
        .agg(n=("relative_coordinate", "size"),
             cmin=("relative_coordinate", "min"),
             cmax=("relative_coordinate", "max"),
             events=("event_count", "sum"),
             zero_frac=("event_count", lambda s: float((s == 0).mean())))
        .reset_index()
        .to_dict("records")
    )

    # ---------- 4. all workbook image sheets ----------
    import openpyxl

    wb = openpyxl.load_workbook(
        DATA / "source_data" / "all_source_data_deidentified.xlsx",
        read_only=True,
        data_only=True,
    )
    img_sheets = [s for s in wb.sheetnames if wb[s].max_col > 100]
    meta_rows = {}
    for s in img_sheets:
        ws = wb[s]
        hdr = [c.value for c in next(ws.iter_rows(min_row=3, max_row=3))]
        bnd = [c.value for c in next(ws.iter_rows(min_row=4, max_row=4))]
        meta_rows[s] = {
            "header_cols_C_G": [hdr[2:7]],
            "bounds_row": [bnd[1:7]],
            "sheet_title_cell": next(ws.iter_rows(min_row=1, max_row=1)).__next__().value,
        }
    wb.close()
    print("\n=== image sheets metadata ===")
    for s, m in meta_rows.items():
        print(s, m["bounds_row"], "| header:", m["header_cols_C_G"][0])
    report["image_sheet_meta"] = meta_rows

    # compare workbook grid vs normalized view for the 5 events
    sheet_for_event = {
        "event_1": "Sheet_003", "event_2": "Sheet_004", "event_3": "Sheet_005",
        "event_4": "Sheet_006", "event_5": "Sheet_007",
    }
    print("\n=== workbook grid vs normalized view ===")
    cmp_rows = []
    for eid, sh in sheet_for_event.items():
        g = load_workbook_grid(sh)
        n = ev[eid]["img"]
        same = np.allclose(g, n, equal_nan=True)
        # also the transposed orientation
        sameT = g.shape == n.T.shape and np.allclose(g, n.T, equal_nan=True)
        print(f"{eid} {sh}: grid{g.shape} identical={same} identical_transpose={sameT} "
              f"maxdiff={np.nanmax(np.abs(g - n)):.3e}")
        cmp_rows.append({"event": eid, "sheet": sh, "identical": bool(same),
                         "identical_transpose": bool(sameT),
                         "max_abs_diff": float(np.nanmax(np.abs(g - n)))})
    report["workbook_vs_normalized"] = cmp_rows

    (OUT / "s02_exploration.json").write_text(json.dumps(report, indent=2, default=str) + "\n")
    np.savez_compressed(
        OUT / "s02_event_grids.npz",
        **{f"{k}__img": v["img"] for k, v in ev.items()},
        **{f"{k}__freq": v["freq"] for k, v in ev.items()},
        **{f"{k}__times": v["times"] for k, v in ev.items()},
    )
    print("\nwrote", OUT / "s02_exploration.json")


if __name__ == "__main__":
    main()