#!/usr/bin/env python3
"""S01 - Complete inventory of the task dataset.

Enumerates every sheet of the deidentified source workbook and every normalized
view, recovering for each: anonymous label(s), header row, data extent, and a
coarse structural classification.  Writes machine-readable inventory artifacts.

No scientific conclusions are drawn here; this establishes the analysis units
and the missingness boundary required by the task contract.
"""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / "data"
OUT = ROOT / "output" / "results"
OUT.mkdir(parents=True, exist_ok=True)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def classify(rows: int, cols: int, headers: list, labels: list) -> str:
    if rows <= 1 and cols <= 1:
        return "empty_placeholder"
    if cols > 100:
        return "dynamic_spectrum_image"
    if "cml angle" in [str(h).strip().lower() for h in headers]:
        return "angle_time_series"
    if len(labels) > 1:
        return "multi_block_distribution"
    if rows > 200 and cols == 4:
        return "count_series_image_projection"
    return "single_block_distribution"


def main() -> None:
    file_rows = []
    for p in sorted(DATA.rglob("*")):
        if p.is_file():
            file_rows.append(
                {
                    "path": p.relative_to(ROOT).as_posix(),
                    "bytes": p.stat().st_size,
                    "sha256": sha256(p),
                }
            )
    (OUT / "s01_file_hashes.json").write_text(
        json.dumps(file_rows, indent=2) + "\n", encoding="utf-8"
    )

    wb = openpyxl.load_workbook(
        DATA / "source_data" / "all_source_data_deidentified.xlsx",
        read_only=True,
        data_only=True,
    )
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        max_row, max_col = ws.max_row, ws.max_column
        head = []
        labels = []
        n_rows_scanned = 0
        for i, row in enumerate(ws.iter_rows(max_row=min(max_row, 8), values_only=True)):
            n_rows_scanned = i
            vals = list(row)
            if i == 2:
                head = [v for v in vals[2:10] if v is not None]
            if i >= 3 and len(vals) > 1 and vals[1] is not None:
                labels.append(str(vals[1]).strip())
        # full column-B label census (cheap: column B only)
        col_b = set()
        for row in ws.iter_rows(min_col=2, max_col=2, values_only=True):
            v = row[0]
            if v is not None and isinstance(v, str):
                col_b.add(v.strip())
        sheets.append(
            {
                "sheet": name,
                "max_row": max_row,
                "max_col": max_col,
                "headers": [str(h) for h in head],
                "labels_colB": sorted(col_b),
                "class": classify(max_row, max_col, head, sorted(col_b)),
            }
        )
    wb.close()
    (OUT / "s01_workbook_inventory.json").write_text(
        json.dumps(sheets, indent=2) + "\n", encoding="utf-8"
    )

    # normalized views
    views = {}
    meta = pd.read_csv(DATA / "normalized_views" / "dynamic_spectrum_event_metadata.tsv", sep="\t")
    views["dynamic_spectrum_event_metadata"] = {
        "n_rows": len(meta),
        "columns": list(meta.columns),
        "event_ids": meta["event_id"].tolist(),
    }
    pop = pd.read_csv(DATA / "normalized_views" / "drift_population_series.tsv", sep="\t")
    views["drift_population_series"] = {
        "n_rows": len(pop),
        "columns": list(pop.columns),
        "n_series": int(pop["population_series_id"].nunique()),
        "series": sorted(pop["population_series_id"].unique().tolist()),
        "groups": sorted(pop["anonymous_source_group"].unique().tolist()),
        "channels": sorted(pop["polarization_channel"].unique().tolist()),
        "regimes": sorted(pop["drift_regime"].unique().tolist()),
        "total_events": float(pop["event_count"].sum()),
    }
    ori = pd.read_csv(DATA / "normalized_views" / "orientation_count_series.tsv", sep="\t")
    views["orientation_count_series"] = {
        "n_rows": len(ori),
        "columns": list(ori.columns),
        "n_series": int(ori["orientation_series_id"].nunique()),
        "series": sorted(ori["orientation_series_id"].unique().tolist()),
        "total_events": float(ori["event_count"].sum()),
    }
    ds_files = sorted(DATA.glob("normalized_views/event_*_dynamic_spectrum.tsv"))
    views["dynamic_spectra"] = []
    for f in ds_files:
        d = pd.read_csv(f, sep="\t", nrows=2)
        views["dynamic_spectra"].append(
            {"file": f.name, "n_freq_rows": None, "n_time_cols": d.shape[1] - 1}
        )
    (OUT / "s01_normalized_view_inventory.json").write_text(
        json.dumps(views, indent=2) + "\n", encoding="utf-8"
    )

    print(json.dumps({"sheets": len(sheets), "files": len(file_rows)}, indent=2))
    for s in sheets:
        print(
            f"{s['sheet']:>11} {s['max_row']:>6}x{s['max_col']:<4} {s['class']:<34} "
            f"labels={s['labels_colB'][:3]}"
        )
    print("\npop series:", views["drift_population_series"]["n_series"], views["drift_population_series"]["total_events"])
    print("orientation series:", views["orientation_count_series"])


if __name__ == "__main__":
    main()