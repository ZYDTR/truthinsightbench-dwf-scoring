#!/usr/bin/env python3
"""S10 (dimension: provenance & coverage) - exploratory dump of raw workbook layouts.

Prints the first rows of representative sheets of each family so that the exact
row/column conventions used by the tidy parsers (s11/s12) are established from
the file itself and not from assumptions.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "data" / "source_data" / "all_source_data_deidentified.xlsx"

PROBE = ["Sheet_001", "Sheet_003", "Sheet_008", "Sheet_009", "Sheet_019",
         "Sheet_020", "Sheet_024", "Sheet_025", "Sheet_039", "Sheet_045",
         "Sheet_046"]


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    print("sheetnames:", len(wb.sheetnames), wb.sheetnames)
    for name in PROBE:
        ws = wb[name]
        print(f"\n===== {name}  (max_row={ws.max_row}, max_col={ws.max_column}) =====")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, max_col=7,
                                             values_only=True)):
            print(f"  r{i + 1}: {row}")
    wb.close()

    # row-type fingerprint of one image sheet: where the numeric 425x425 block sits
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["Sheet_003"]
    kinds = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=440, max_col=6,
                                         values_only=True)):
        a, b = row[0], row[1]
        numeric_tail = sum(1 for v in row[1:6] if isinstance(v, (int, float)))
        kinds.append((i + 1, repr(a)[:28], repr(b)[:28], numeric_tail))
    for k in kinds[:12] + kinds[-12:]:
        print("Sheet_003 row", k)
    wb.close()


if __name__ == "__main__":
    main()